#!/usr/bin/env python3
"""Generate a Q&A-only document for easy copy-paste into submission forms."""
import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MD_PATH = r"E:\重要内容\GenAI比赛\Inter-university 比赛\voltguard-site\docs\INTERIM_REPORT_FINAL.md"
OUTPUT_DOCX = r"E:\重要内容\GenAI比赛\Inter-university 比赛\INTERIM_REPORT_QA.docx"

def set_cell_shading(cell, color_hex):
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color_hex)
    cell._tc.get_or_add_tcPr().append(shading)

def add_qa_block(doc, q_num, q_text, answer_lines):
    """Add a question-answer block with clear visual separation."""
    # Question header
    q_para = doc.add_paragraph()
    q_para.paragraph_format.space_before = Pt(16)
    q_para.paragraph_format.space_after = Pt(6)

    q_run = q_para.add_run(f"Q{q_num}. {q_text}")
    q_run.bold = True
    q_run.font.size = Pt(12)
    q_run.font.color.rgb = RGBColor(0x0c, 0x2f, 0x3d)

    # Separator line
    sep = doc.add_paragraph()
    sep.paragraph_format.space_before = Pt(2)
    sep.paragraph_format.space_after = Pt(8)
    sep_run = sep.add_run("─" * 60)
    sep_run.font.color.rgb = RGBColor(0xd8, 0xe1, 0xe3)
    sep_run.font.size = Pt(8)

    # Answer content
    for line in answer_lines:
        line = line.strip()
        if not line:
            continue

        # Parse inline markdown
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(4)
        para.paragraph_format.line_spacing = 1.3

        # Handle bullet points
        if line.startswith('- '):
            line = line[2:]
            para.paragraph_format.left_indent = Inches(0.2)
            bullet_run = para.add_run("• ")
            bullet_run.font.size = Pt(10.5)

        # Parse **bold** and *italic*
        parts = re.split(r'(\*\*.*?\*\*|\*[^*]+?\*)', line)
        for part in parts:
            if not part:
                continue
            if part.startswith('**') and part.endswith('**') and len(part) > 4:
                run = para.add_run(part[2:-2])
                run.bold = True
                run.font.color.rgb = RGBColor(0x0c, 0x2f, 0x3d)
            elif part.startswith('*') and part.endswith('*') and len(part) > 2 and not part.startswith('**'):
                run = para.add_run(part[1:-1])
                run.italic = True
            else:
                para.add_run(part)

        for run in para.runs:
            run.font.size = Pt(10.5)

    # Bottom spacing
    doc.add_paragraph()

def extract_qa_from_md(md_text):
    """Extract Q1-Q23 answers from the markdown report."""
    qa_map = {}

    # Define question mapping: (number, short_title, full_question)
    questions = [
        (1, "Problem", "What specific problem or pain point is your project addressing?"),
        (2, "Research", "What research have you conducted to understand the problem better? (Reference supporting studies, statistics, or data collected by your team)"),
        (3, "Significance", "Why is this problem significant? Who does it affect?"),
        (4, "Existing Solutions", "What existing solutions are available and how do they fall short?"),
        (5, "Objectives", "Describe the primary objectives of your project."),
        (6, "Boundaries", "What are the boundaries of your project (what will the project not address)?"),
        (7, "Assumptions", "What assumptions are you making about the problem/pain point and its context?"),
        (8, "Target Audience", "Who is your target audience or user base?"),
        (9, "Client", "Who is your client (the group paying for the solution)?"),
        (10, "Needs Gathering", "How did you gather information about the needs and preferences of your target audience? (Including testing your solution)"),
        (11, "Team Structure", "How is your team structured, and what roles does each member play?"),
        (12, "Communication", "What communication strategies are you using to ensure effective collaboration?"),
        (13, "Achievements", "What have you achieved so far?"),
        (14, "Current Work", "What part of your project are you currently working on?"),
        (15, "Next Steps", "What are your immediate next steps or milestones?"),
        (16, "Success Evaluation", "How will you evaluate the success of your project as it progresses?"),
        (17, "Challenges", "What challenges or obstacles have you encountered so far? How did you address these challenges?"),
        (18, "Future Challenges", "Do you foresee any potential challenges in the upcoming stages of your project? How do you intend to overcome these challenges?"),
        (19, "Resources", "What resources (tools, data, libraries, etc.) have you utilised so far in your project? How have these resources contributed to your progress?"),
        (20, "Missing Resources", "Are there any resources (hardware, software, data, expertise) that you need but currently lack?"),
        (21, "Learnings", "What have you learned so far, either about the problem you're addressing, the technologies you're using, or working as a team?"),
        (22, "Evolution", "How has your understanding of the problem or solution evolved since you started working on it?"),
        (23, "Other Comments", "What other issues or comments does your team want to raise?"),
    ]

    lines = md_text.split('\n')
    current_q = None
    current_ans = []

    for i, line in enumerate(lines):
        # Detect question headers
        if line.startswith('### Q'):
            # Save previous Q
            if current_q and current_ans:
                qa_map[current_q] = current_ans

            # Extract Q number
            match = re.match(r'### Q(\d+)\.', line)
            if match:
                current_q = int(match.group(1))
                current_ans = []
            continue

        # Detect section boundaries
        if line.startswith('---') or line.startswith('## Part'):
            if current_q and current_ans:
                qa_map[current_q] = current_ans
                current_q = None
                current_ans = []
            continue

        # Skip table separator lines
        if line.strip() and all(c in '-|: ' for c in line.strip()):
            continue

        # Collect answer lines
        if current_q is not None:
            # Skip heading lines within answers
            if line.startswith('#'):
                continue
            # Skip question mapping table
            if '| Question |' in line or '| Q1 ' in line or '|------|' in line:
                continue
            current_ans.append(line)

    # Save last Q
    if current_q and current_ans:
        qa_map[current_q] = current_ans

    return qa_map, questions

def build_qa_doc():
    print("Reading markdown...")
    with open(MD_PATH, 'r', encoding='utf-8') as f:
        md_text = f.read()

    print("Extracting Q&A...")
    qa_map, questions = extract_qa_from_md(md_text)

    print(f"Found answers for {len(qa_map)} questions")

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("VoltGuard AI — Interim Progress Report")
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0c, 0x2f, 0x3d)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Q&A Format — For Online Submission")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x16, 0x8b, 0x78)

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = info.add_run("Team 1023 / LadyBanana, HKBU | 4 October 2026")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x60, 0x74, 0x7d)

    doc.add_paragraph()

    # Instructions
    instr = doc.add_paragraph()
    instr.paragraph_format.space_after = Pt(12)
    run = instr.add_run("Instructions: Copy the answer text below each question and paste into the corresponding field on the submission form.")
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x60, 0x74, 0x7d)

    doc.add_paragraph()

    # Output each Q&A
    for q_num, short_title, full_q in questions:
        ans_lines = qa_map.get(q_num, ["[Answer not found in report]"])
        add_qa_block(doc, q_num, full_q, ans_lines)

    doc.save(OUTPUT_DOCX)
    print(f"Done: {OUTPUT_DOCX}")

if __name__ == '__main__':
    build_qa_doc()
