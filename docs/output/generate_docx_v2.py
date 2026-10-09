#!/usr/bin/env python3
"""Generate final DOCX with proper markdown bold/italic handling."""
import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MD_PATH = r"E:\重要内容\GenAI比赛\Inter-university 比赛\voltguard-site\docs\INTERIM_REPORT_FINAL.md"
CHARTS_DIR = r"E:\重要内容\GenAI比赛\Inter-university 比赛\voltguard-site\docs\output\charts"
OUTPUT_DOCX = r"E:\重要内容\GenAI比赛\Inter-university 比赛\INTERIM_REPORT.docx"

def set_cell_shading(cell, color_hex):
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color_hex)
    cell._tc.get_or_add_tcPr().append(shading)

def add_formatted_paragraph(doc, md_text, style='Normal', alignment=None, indent=0):
    """Parse inline markdown bold/italic and add to paragraph."""
    para = doc.add_paragraph(style=style)
    if alignment:
        para.alignment = alignment
    if indent > 0:
        para.paragraph_format.left_indent = Inches(indent)

    # Parse inline markdown: **bold**, *italic*, `code`
    # Split by markdown patterns
    parts = re.split(r'(\*\*.*?\*\*|\*[^*]+?\*|\`[^`]+?\`)', md_text)

    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            run = para.add_run(part[2:-2])
            run.bold = True
            run.font.color.rgb = RGBColor(0x0c, 0x2f, 0x3d)
        elif part.startswith('*') and part.endswith('*') and len(part) > 2 and not part.startswith('**'):
            run = para.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            run = para.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x60, 0x74, 0x7d)
        else:
            para.add_run(part)

    for run in para.runs:
        if run.font.size is None:
            run.font.size = Pt(10.5)

    return para

def add_cover_page(doc):
    for _ in range(6):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("VoltGuard AI")
    run.font.size = Pt(36)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0c, 0x2f, 0x3d)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Interim Progress Report")
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0x16, 0x8b, 0x78)

    event = doc.add_paragraph()
    event.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = event.add_run("Hack for SDGs 2026")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x60, 0x74, 0x7d)

    track = doc.add_paragraph()
    track.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = track.add_run("Track 1 — Green Transport & Smart Mobility")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x60, 0x74, 0x7d)

    for _ in range(4):
        doc.add_paragraph()

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = info.add_run("Team 1023 / LadyBanana\nHong Kong Baptist University (HKBU)")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x10, 0x25, 0x2f)

    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date_para.add_run("4 October 2026")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x60, 0x74, 0x7d)

    doc.add_page_break()

def add_toc_page(doc):
    toc_heading = doc.add_heading("Table of Contents", level=1)
    for run in toc_heading.runs:
        run.font.color.rgb = RGBColor(0x0c, 0x2f, 0x3d)

    toc_items = [
        ("Executive Summary", 1), ("SDG Alignment", 1), ("Question Mapping", 1),
        ("Part 1: Problem and Scope", 1),
        ("    Q1. Specific Pain Point", 2), ("    Q2. Research & Data", 2),
        ("    Q3. Significance & Affected Groups", 2), ("    Q4. Existing Solutions & Gaps", 2),
        ("Part 2: Proposed Solution", 1),
        ("    Q5. Primary Objectives", 2), ("    Q6. Project Boundaries", 2),
        ("    Q7. Assumptions", 2), ("    Q8. Target Audience", 2),
        ("    Q9. Client", 2), ("    Q10. Needs Gathering & Testing", 2),
        ("Part 3: Teamwork", 1),
        ("    Q11. Team Structure", 2), ("    Q12. Communication Strategies", 2),
        ("Part 4: Progress", 1),
        ("    Q13. Achievements", 2), ("    Q14. Current Work", 2),
        ("    Q15. Next Steps & Milestones", 2), ("    Q16. Success Evaluation", 2),
        ("Part 5: Challenges", 1),
        ("    Q17. Encountered Challenges", 2), ("    Q18. Upcoming Challenges", 2),
        ("Part 6: Resources", 1),
        ("    Q19. Resources Utilised", 2), ("    Q20. Missing Resources", 2),
        ("Part 7: Learnings", 1),
        ("    Q21. Learnings", 2), ("    Q22. Evolution", 2), ("    Q23. Other Comments", 2),
    ]

    for item, level in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3 * (level - 1))
        run = p.add_run(item)
        run.font.size = Pt(11 if level == 1 else 10)
        run.font.color.rgb = RGBColor(0x10, 0x25, 0x2f) if level == 1 else RGBColor(0x60, 0x74, 0x7d)

    doc.add_page_break()

def flush_table(doc, rows):
    if not rows:
        return
    max_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=max_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, row_data in enumerate(rows):
        row = table.rows[i]
        for j, cell_text in enumerate(row_data):
            if j < max_cols:
                cell = row.cells[j]
                # Strip markdown formatting from table cells
                clean_text = re.sub(r'\*\*(.*?)\*\*', r'\1', cell_text)
                clean_text = re.sub(r'\*(.*?)\*', r'\1', clean_text)
                cell.text = clean_text
                if i == 0:
                    set_cell_shading(cell, '0c2f3d')
                    for para in cell.paragraphs:
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in para.runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            run.font.size = Pt(9)
                else:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(9)
                    if i % 2 == 0:
                        set_cell_shading(cell, 'f4f7f6')
    doc.add_paragraph()

def parse_md_to_docx(md_text, doc):
    lines = md_text.split('\n')
    i = 0
    in_table = False
    table_rows = []

    while i < len(lines):
        line = lines[i].rstrip()

        if not line and i < 5:
            i += 1
            continue

        # Headings
        if line.startswith('# ') and not line.startswith('## '):
            if in_table: flush_table(doc, table_rows); in_table = False; table_rows = []
            h = doc.add_heading(line[2:], level=1)
            for r in h.runs: r.font.color.rgb = RGBColor(0x0c, 0x2f, 0x3d); r.font.size = Pt(18)
            i += 1; continue
        elif line.startswith('## '):
            if in_table: flush_table(doc, table_rows); in_table = False; table_rows = []
            h = doc.add_heading(line[3:], level=2)
            for r in h.runs: r.font.color.rgb = RGBColor(0x0c, 0x2f, 0x3d); r.font.size = Pt(14)
            i += 1; continue
        elif line.startswith('### '):
            if in_table: flush_table(doc, table_rows); in_table = False; table_rows = []
            h = doc.add_heading(line[4:], level=3)
            for r in h.runs: r.font.color.rgb = RGBColor(0x0c, 0x2f, 0x3d); r.font.size = Pt(12)
            i += 1; continue
        elif line.startswith('#### '):
            if in_table: flush_table(doc, table_rows); in_table = False; table_rows = []
            h = doc.add_heading(line[5:], level=4)
            for r in h.runs: r.font.color.rgb = RGBColor(0x0c, 0x2f, 0x3d); r.font.size = Pt(11)
            i += 1; continue

        if line == '---':
            if in_table: flush_table(doc, table_rows); in_table = False; table_rows = []
            if i + 1 < len(lines) and lines[i+1].strip().startswith('# '):
                doc.add_page_break()
            i += 1; continue

        # Table
        if '|' in line and not line.startswith('**'):
            parts = [p.strip() for p in line.split('|')]
            parts = [p for p in parts if p]
            if parts:
                if all(set(p) <= set('-|: ') for p in parts):
                    i += 1; continue
                if not in_table:
                    in_table = True
                    table_rows = []
                table_rows.append(parts)
            i += 1; continue
        elif in_table and (not line or '|' not in line):
            flush_table(doc, table_rows); in_table = False; table_rows = []

        # Bullet points
        if line.startswith('- ') or line.startswith('* '):
            text = line[2:]
            add_formatted_paragraph(doc, text, style='List Bullet', indent=0.25)
            i += 1; continue

        # Numbered lists
        match = re.match(r'^\d+\.\s+(.*)$', line)
        if match:
            if in_table: flush_table(doc, table_rows); in_table = False; table_rows = []
            add_formatted_paragraph(doc, match.group(1), style='List Number', indent=0.25)
            i += 1; continue

        # Blockquote
        if line.startswith('>'):
            if in_table: flush_table(doc, table_rows); in_table = False; table_rows = []
            p = doc.add_paragraph(line[1:].strip())
            p.paragraph_format.left_indent = Inches(0.3)
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x60, 0x74, 0x7d)
                r.font.size = Pt(10.5)
            i += 1; continue

        # Regular paragraph with inline formatting
        if line.strip():
            if in_table: flush_table(doc, table_rows); in_table = False; table_rows = []
            add_formatted_paragraph(doc, line)
            doc.paragraphs[-1].paragraph_format.space_after = Pt(4)

        i += 1

    if in_table:
        flush_table(doc, table_rows)

def add_header_footer(doc):
    for section in doc.sections:
        # Header
        header = section.header
        header.is_linked_to_previous = False
        para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = para.add_run("VoltGuard AI — Interim Progress Report")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x60, 0x74, 0x7d)
        run.italic = True

        # Footer
        footer = section.footer
        footer.is_linked_to_previous = False
        para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run("Team 1023 / LadyBanana, HKBU  |  ")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x60, 0x74, 0x7d)

def add_charts(doc):
    doc.add_page_break()
    h = doc.add_heading('Appendix: Data Visualisations', level=1)
    for r in h.runs: r.font.color.rgb = RGBColor(0x0c, 0x2f, 0x3d)

    charts = [
        (os.path.join(CHARTS_DIR, 'chart_stakeholders.png'), 'Figure 1: Stakeholder Impact & Urgency Assessment'),
        (os.path.join(CHARTS_DIR, 'chart_nasa_data.png'), 'Figure 2: NASA Ames PCoE Dataset Statistics'),
        (os.path.join(CHARTS_DIR, 'chart_ev_growth.png'), 'Figure 3: Hong Kong EV Registration Growth (2022-2026)'),
        (os.path.join(CHARTS_DIR, 'chart_architecture.png'), 'Figure 4: VoltGuard Two-Tier Safety Architecture'),
        (os.path.join(CHARTS_DIR, 'chart_milestones.png'), 'Figure 5: Project Milestones Timeline (Oct 2026)'),
    ]

    for img_path, caption in charts:
        if os.path.exists(img_path):
            doc.add_paragraph()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(img_path, width=Inches(5.8))

            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            crun = cap.add_run(caption)
            crun.italic = True
            crun.font.size = Pt(9)
            crun.font.color.rgb = RGBColor(0x60, 0x74, 0x7d)
            doc.add_paragraph()

def main():
    print("Reading markdown...")
    with open(MD_PATH, 'r', encoding='utf-8') as f:
        md_text = f.read()

    print("Creating document...")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)

    add_cover_page(doc)
    add_toc_page(doc)
    parse_md_to_docx(md_text, doc)
    add_charts(doc)
    add_header_footer(doc)

    doc.save(OUTPUT_DOCX)
    print(f"Done: {OUTPUT_DOCX}")

if __name__ == '__main__':
    main()
